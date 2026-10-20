"""Create sample_data/timetable_input_sample.xlsx.

The file follows the structure the app expects (sheets Classes, Subjects,
Teachers, Settings) and contains example data that passes validation and
produces a timetable. It is what the "Download sample Excel" button serves.

Run from the project folder:   python tools/create_sample_excel.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = (
    Path(__file__).resolve().parent.parent
    / "sample_data"
    / "timetable_input_sample.xlsx"
)

HEADER_FILL = PatternFill("solid", fgColor="2F4B8F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
AUTO_FILL = PatternFill("solid", fgColor="E8EDF7")
NOTE_FONT = Font(italic=True, color="666666")
THIN = Side(style="thin", color="D0D7E5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MAX_ROWS = 300  # rows covered by dropdowns / auto-filled class IDs

CLASSES = [
    # course_name, branch_name, semester, section
    ("B.Tech", "CSE", 1, "A"),
    ("B.Tech", "CSE", 3, "A"),
    ("B.Tech", "Civil", 1, "A"),
    ("BCA", "General", 1, "A"),
]

SUBJECTS = [
    # class_id, subject_name, weekly_periods, periods_per_session, subject_type
    ("B.Tech CSE 1A", "Mathematics-I", 4, 1, "Theory"),
    ("B.Tech CSE 1A", "Physics", 4, 1, "Theory"),
    ("B.Tech CSE 1A", "Programming in C", 4, 1, "Theory"),
    ("B.Tech CSE 1A", "English", 3, 1, "Theory"),
    ("B.Tech CSE 1A", "Environmental Science", 2, 1, "Theory"),
    ("B.Tech CSE 1A", "Physics Lab", 2, 2, "Lab"),
    ("B.Tech CSE 1A", "Programming Lab", 4, 2, "Lab"),

    ("B.Tech CSE 3A", "Data Structures", 5, 1, "Theory"),
    ("B.Tech CSE 3A", "Computer Organization", 4, 1, "Theory"),
    ("B.Tech CSE 3A", "Object Oriented Programming", 4, 1, "Theory"),
    ("B.Tech CSE 3A", "Discrete Mathematics", 4, 1, "Theory"),
    ("B.Tech CSE 3A", "Professional Communication", 2, 1, "Theory"),
    ("B.Tech CSE 3A", "DS Lab", 4, 2, "Lab"),

    ("B.Tech Civil 1A", "Mathematics-I", 4, 1, "Theory"),
    ("B.Tech Civil 1A", "Physics", 4, 1, "Theory"),
    ("B.Tech Civil 1A", "Surveying", 3, 1, "Theory"),
    ("B.Tech Civil 1A", "English", 3, 1, "Theory"),
    ("B.Tech Civil 1A", "Environmental Science", 2, 1, "Theory"),
    ("B.Tech Civil 1A", "Engineering Drawing", 4, 2, "Lab"),

    ("BCA General 1A", "Computer Fundamentals", 4, 1, "Theory"),
    ("BCA General 1A", "Mathematics-I", 3, 1, "Theory"),
    ("BCA General 1A", "English", 3, 1, "Theory"),
    ("BCA General 1A", "Programming in C", 4, 1, "Theory"),
    ("BCA General 1A", "Programming Lab", 2, 2, "Lab"),
]

TEACHERS = [
    # teacher_name, subjects (comma separated)
    ("Amit", "Mathematics-I, Discrete Mathematics"),
    ("Ravi", "Physics, Physics Lab"),
    ("Indu", "Programming in C, Programming Lab"),
    ("Rahul", "Data Structures, DS Lab, Computer Organization, Mini Project"),
    ("Neha", "Object Oriented Programming, Computer Fundamentals, Skill Development"),
    ("Alpana", "English, Professional Communication, Skill Development"),
    ("Vikram", "English, Environmental Science, Technical Seminar"),
    ("Sohan Lal", "Engineering Drawing, Surveying"),
]

COMMON_CLASSES = [
    # group_name, class_ids, subject_name, weekly_periods,
    # periods_per_session, mode, fixed_day, fixed_period
    ("Technical Seminar", "B.Tech CSE 1A, B.Tech Civil 1A",
     "Technical Seminar", 1, 1, "joint", "Saturday", 6),
    ("Mini Project", "B.Tech CSE 1A, B.Tech CSE 3A",
     "Mini Project", 2, 2, "joint", None, None),
    ("Skill Development", "B.Tech CSE 3A, BCA General 1A",
     "Skill Development", 2, 1, "parallel", None, None),
]

SETTINGS = [
    ("working_days", "Monday,Tuesday,Wednesday,Thursday,Friday,Saturday"),
    ("college_start_time", "09:00"),
    ("college_end_time", "16:30"),
    ("period_duration_minutes", 50),
    ("lunch_start_time", "12:20"),
    ("lunch_duration_minutes", 40),
    ("compact_timetable", "preferred"),
]


def style_header(ws, widths):
    for col, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def add_rows(ws, rows, text_columns=()):
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            if c in text_columns:
                cell.number_format = "@"


def build_instructions(ws):
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95

    def title(row, text):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, size=14, color="2F4B8F")

    def line(row, label, text):
        a = ws.cell(row=row, column=1, value=label)
        a.font = Font(bold=True)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=row, column=2, value=text)
        b.alignment = Alignment(wrap_text=True, vertical="top")

    r = 1
    title(r, "How to use this file"); r += 2
    line(r, "Step 1", "Replace the example data in the Classes, Subjects, Teachers and Settings sheets with your own. Keep the header row (row 1) exactly as it is."); r += 1
    line(r, "Step 2", "Delete the example rows you do not need. Blank rows are ignored."); r += 1
    line(r, "Step 3", "Save the file as .xlsx and upload it on the timetable page."); r += 2

    title(r, "Sheet: Classes"); r += 1
    line(r, "course_name", "Course, e.g. B.Tech, BCA, Diploma."); r += 1
    line(r, "branch_name", "Branch or stream, e.g. CSE, Civil, General."); r += 1
    line(r, "semester", "Whole number, e.g. 1, 3, 5."); r += 1
    line(r, "section", "Section letter, e.g. A, B."); r += 1
    line(r, "class_id (grey column)", "Filled in automatically. It is the name of the class used by the Subjects sheet: course + branch + semester + section, e.g. 'B.Tech CSE 3A'. Do not type in this column."); r += 2

    title(r, "Sheet: Subjects (one row per subject per class)"); r += 1
    line(r, "class_id", "Pick the class from the drop-down list. If you type it yourself it must match the grey class_id column in the Classes sheet."); r += 1
    line(r, "subject_name", "Subject name. It must also appear in the 'subjects' cell of at least one teacher (spelling must match; upper/lower case does not matter)."); r += 1
    line(r, "weekly_periods", "How many periods per week this class has this subject (whole number)."); r += 1
    line(r, "periods_per_session", "Periods in a row for one session. Use 1 for normal lectures and 2 (or more) for labs. weekly_periods must be divisible by this number. Sessions never run across the lunch break."); r += 1
    line(r, "subject_type", "Theory, Lab or Project (information only)."); r += 2

    title(r, "Sheet: Teachers (one row per teacher)"); r += 1
    line(r, "teacher_name", "Each teacher once. Names must be unique."); r += 1
    line(r, "subjects", "All subjects this teacher can teach, separated by commas. If two teachers list the same subject, the system picks one of them for each class, and that teacher teaches the subject to that class all week."); r += 2

    title(r, "Sheet: Common_Classes (optional - delete the rows if you have none)"); r += 1
    line(r, "group_name", "Name of the shared class, e.g. Technical Seminar, Skill Development, Project."); r += 1
    line(r, "class_ids", "Two or more classes that have it at the same time, separated by commas. Copy them from the grey class_id column of the Classes sheet."); r += 1
    line(r, "subject_name", "Must also appear in the 'subjects' cell of at least one teacher. Do NOT list the same subject for these classes in the Subjects sheet."); r += 1
    line(r, "weekly_periods / periods_per_session", "Same meaning as in the Subjects sheet."); r += 1
    line(r, "mode", "joint = ONE teacher teaches all the classes together (one lesson for everyone). parallel = same time slot, but each class has its OWN teacher (needs at least as many teachers for the subject as classes). Empty = joint."); r += 1
    line(r, "fixed_day (optional)", "Force a day, e.g. Saturday. Leave empty to let the system choose."); r += 1
    line(r, "fixed_period (optional)", "Force a period number (1 = first period of the day). Leave empty to let the system choose."); r += 2

    title(r, "Sheet: Settings (keep the names in the 'setting' column)"); r += 1
    line(r, "working_days", "Comma separated days, e.g. Monday,Tuesday,Wednesday,Thursday,Friday,Saturday."); r += 1
    line(r, "college_start_time", "First period starts, 24-hour time such as 09:00."); r += 1
    line(r, "college_end_time", "Last period must end by this time, e.g. 16:30."); r += 1
    line(r, "period_duration_minutes", "Length of one period in minutes, e.g. 50."); r += 1
    line(r, "lunch_start_time", "Lunch begins, e.g. 12:20. Periods are not scheduled during lunch."); r += 1
    line(r, "lunch_duration_minutes", "Lunch length in minutes (0 for no lunch break)."); r += 1
    line(r, "compact_timetable (optional)", "preferred (default) = avoid free periods, such as an empty first period, where possible. strict = lessons must start at period 1 with no free period in between (may make a timetable impossible). off = ignore free periods."); r += 2

    title(r, "Rules the generator follows"); r += 1
    for text in [
        "A class has only one lesson at a time, and a teacher teaches only one class at a time.",
        "A subject has at most one session per day for a class.",
        "One teacher teaches a given subject to a given class for the whole week.",
        "Common classes are scheduled once and block every class that shares them.",
        "Free periods are avoided: lessons are packed from the first period of the day (see compact_timetable).",
        "Multi-period sessions (labs) never start so late that they run over lunch or past the end of the day.",
        "The number of periods per day comes from Settings (start, end, period length and lunch) - it is not fixed.",
    ]:
        ws.cell(row=r, column=1, value="•").alignment = Alignment(horizontal="right")
        cell = ws.cell(row=r, column=2, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1


def main():
    wb = Workbook()

    # Instructions
    ws = wb.active
    ws.title = "Instructions"
    build_instructions(ws)

    # Classes
    ws = wb.create_sheet("Classes")
    ws.append(["course_name", "branch_name", "semester", "section", "class_id"])
    style_header(ws, [16, 16, 12, 12, 26])
    add_rows(ws, CLASSES)
    for r in range(2, MAX_ROWS + 1):
        cell = ws.cell(
            row=r,
            column=5,
            value=(
                f'=IF(A{r}="","",TRIM(A{r})&" "&TRIM(B{r})&" "&'
                f'TRIM(C{r})&TRIM(D{r}))'
            ),
        )
        cell.fill = AUTO_FILL
        cell.border = BORDER
    ws["E1"].fill = PatternFill("solid", fgColor="6B7A99")

    # Subjects
    ws = wb.create_sheet("Subjects")
    ws.append([
        "class_id", "subject_name", "weekly_periods",
        "periods_per_session", "subject_type",
    ])
    style_header(ws, [22, 32, 16, 20, 16])
    add_rows(ws, SUBJECTS)

    class_dropdown = DataValidation(
        type="list",
        formula1=f"=Classes!$E$2:$E${MAX_ROWS}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Unknown class",
        error="Pick a class that exists in the Classes sheet.",
    )
    type_dropdown = DataValidation(
        type="list", formula1='"Theory,Lab,Project"', allow_blank=True
    )
    whole_numbers = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="1",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Whole number needed",
        error="Enter a whole number of 1 or more.",
    )
    ws.add_data_validation(class_dropdown)
    ws.add_data_validation(type_dropdown)
    ws.add_data_validation(whole_numbers)
    class_dropdown.add(f"A2:A{MAX_ROWS}")
    whole_numbers.add(f"C2:D{MAX_ROWS}")
    type_dropdown.add(f"E2:E{MAX_ROWS}")

    # Teachers
    ws = wb.create_sheet("Teachers")
    ws.append(["teacher_name", "subjects"])
    style_header(ws, [20, 70])
    add_rows(ws, TEACHERS)

    # Common_Classes (optional)
    ws = wb.create_sheet("Common_Classes")
    ws.append([
        "group_name", "class_ids", "subject_name", "weekly_periods",
        "periods_per_session", "mode", "fixed_day", "fixed_period",
    ])
    style_header(ws, [22, 44, 24, 16, 20, 12, 14, 14])
    add_rows(ws, COMMON_CLASSES)
    ws.add_data_validation(whole_numbers_common := DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="1",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Whole number needed",
        error="Enter a whole number of 1 or more.",
    ))
    whole_numbers_common.add(f"D2:D{MAX_ROWS}")
    whole_numbers_common.add(f"H2:H{MAX_ROWS}")
    mode_dropdown = DataValidation(
        type="list", formula1='"joint,parallel"', allow_blank=True
    )
    day_dropdown = DataValidation(
        type="list",
        formula1='"Monday,Tuesday,Wednesday,Thursday,Friday,Saturday,Sunday"',
        allow_blank=True,
    )
    ws.add_data_validation(mode_dropdown)
    ws.add_data_validation(day_dropdown)
    mode_dropdown.add(f"F2:F{MAX_ROWS}")
    day_dropdown.add(f"G2:G{MAX_ROWS}")

    # Settings (times stored as text so Excel does not reformat them)
    ws = wb.create_sheet("Settings")
    ws.append(["setting", "value"])
    style_header(ws, [28, 52])
    add_rows(ws, SETTINGS, text_columns=(2,))
    for r in range(2, len(SETTINGS) + 2):
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")

    OUTPUT.parent.mkdir(exist_ok=True)
    wb.save(OUTPUT)
    print("Created", OUTPUT)


if __name__ == "__main__":
    main()
